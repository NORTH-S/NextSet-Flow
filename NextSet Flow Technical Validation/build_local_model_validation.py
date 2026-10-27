from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
MODEL = PROJECT / 'NextSet Flow Intent Model'
METRICS = json.loads((MODEL / 'reports/metrics.json').read_text(encoding='utf-8'))
SELECTION = json.loads((MODEL / 'model/selection.json').read_text(encoding='utf-8'))
OUTPUT = ROOT / 'NextSet Flow Local Model Technical Validation.docx'
BLACK, NAVY, PALE, WHITE, BORDER = '000000', '17324D', 'F1F5F8', 'FFFFFF', 'D9D9D9'

def p(doc, text, style=None, size=None):
    paragraph=doc.add_paragraph(text,style)
    if size:
        for run in paragraph.runs: run.font.size=Pt(size)
    return paragraph

def h(doc,text):
    return p(doc,text,'Heading 1')

def page(doc,title):
    doc.add_page_break()
    h(doc,title)

def cell_format(cell,fill):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=cell._tc.get_or_add_tcPr()
    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),fill);tcpr.append(shading)
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
        edge=OxmlElement('w:'+side)
        for key,value in {'val':'single','sz':'4','color':BORDER}.items():edge.set(qn('w:'+key),value)
        borders.append(edge)
    tcpr.append(borders)
    margins=OxmlElement('w:tcMar')
    for side,value in [('top','80'),('bottom','80'),('start','100'),('end','100')]:
        element=OxmlElement('w:'+side);element.set(qn('w:w'),value);element.set(qn('w:type'),'dxa');margins.append(element)
    tcpr.append(margins)

def table(doc,heads,rows,widths,numeric=()):
    t=doc.add_table(rows=1,cols=len(heads));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for row_index,values in enumerate([heads]+rows):
        row=t.rows[0] if row_index==0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if row_index==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i,value in enumerate(values):
            c=row.cells[i];c.width=Inches(widths[i]);cell_format(c,NAVY if not row_index else PALE if row_index%2 else WHITE)
            para=c.paragraphs[0];para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
            if i in numeric:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=para.add_run(str(value));run.font.name='Arial';run.font.size=Pt(9.5);run.font.bold=row_index==0;run.font.color.rgb=RGBColor.from_string(WHITE if not row_index else BLACK)
    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=1;spacer.add_run().font.size=Pt(2)
    return t

def link(doc,label,url):
    paragraph=doc.add_paragraph()
    paragraph.paragraph_format.space_after=Pt(5)
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),rel)
    run=OxmlElement('w:r');rpr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),BLACK);rpr.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');rpr.append(size);run.append(rpr)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyperlink.append(run);paragraph._p.append(hyperlink)
    return paragraph

def make():
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
    s.top_margin=Inches(.72);s.bottom_margin=Inches(.7);s.left_margin=Inches(.78);s.right_margin=Inches(.78)
    s.header_distance=Inches(.28);s.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string(BLACK)
        st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial');st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
        for border in list(st._element.xpath('.//w:pBdr')):border.getparent().remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.1;normal.paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(25);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(7)
    d.styles['Subtitle'].font.size=Pt(13);d.styles['Subtitle'].paragraph_format.space_after=Pt(9)
    d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(13);d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    d.styles['Heading 2'].font.size=Pt(11.5);d.styles['Heading 2'].font.bold=True
    for name in ['Heading 1','Heading 2']:
        d.styles[name].paragraph_format.keep_with_next=True
    d.core_properties.title='NextSet Flow Local Model Technical Validation'
    d.core_properties.subject='Validation of a trained local workout focus classifier and its demo integration'
    d.core_properties.author='The Coalition'
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('NextSet Flow  |  Technical validation  |  ');run.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

    p(d,'NextSet Flow Technical Validation','Title')
    p(d,'Local Workout Classification and Demo Integration','Subtitle')
    p(d,'Enterprise Roll Out  |  The Coalition  |  27 October 2026',size=9)
    p(d,'We have trained and integrated a small bilingual classifier that recognises workout focus without calling an external model API. The current version is suitable for a course demonstration and further testing. It is not yet validated for dependable use with real gym members. On 105 separately authored synthetic requests, raw classification accuracy was 72.4%; the integrated software passed 78 functional checks. These are different forms of evidence and should not be treated as interchangeable.')
    h(d,'Purpose and scope')
    p(d,'The technical question is whether a user can describe a workout in English, Chinese or a mixture of both, obtain a supported training focus, and pass a validated request to the existing gym scheduler. The classifier supports seven labels: legs, upper body, full body, push, pull, clarification required and unsupported request. It classifies text; it does not generate exercise prescriptions or predict how long equipment will remain occupied.')
    h(d,'System responsibilities')
    table(d,['Component','Implementation','Responsibility'],[
        ['Workout focus','Trained local classifier','Recognise a supported focus or request clarification.'],
        ['Request fields','Explicit parsing rules','Extract duration, experience, exclusions and preferences.'],
        ['Session planning','Deterministic scheduler','Allocate suitable stations and update routes and timings.'],
    ],[1.30,1.65,3.72])
    p(d,'Keeping these responsibilities separate allows the language component to remain small while equipment decisions use explicit state and exercise relationships. An estimate passing its expected finish time does not make a station free; the scheduling logic requires a finish confirmation. Later session steps remain forecasts rather than reservations.')
    h(d,'Model design')
    p(d,'The model uses character sequences of two to five characters, weighted with term frequency and inverse document frequency, or TF-IDF [1]. A seven-class L2 logistic regression model converts these features into class scores [2]. The selected model uses 12,000 character features and regularisation setting C = 4. The export contains 84,007 learned weight and bias values and is approximately 1.36 MiB. No pretrained language-model weights were used.')

    page(d,'Training data and evaluation method')
    p(d,'The development corpus was created specifically for this experiment. Every request is assistant-authored synthetic text, not a request collected from a gym member. The examples cover English, Chinese and mixed-language wording, including corrections, exclusions, ambiguous goals and requests outside the supported task. The label boundaries are product choices; for example, a general back request is upper body, whereas an explicit pull-day or back-and-biceps request is pull.')
    table(d,['Dataset','Size','Use'],[
        ['Training seeds','350 requests','50 requests per class; used to learn features and weights.'],
        ['Augmented training','1,050 rows','Training-only additions of time or queue context. These are not 1,050 independent users.'],
        ['Validation','98 requests','14 per class; used to select model settings and uncertainty thresholds.'],
        ['Frozen final test','105 requests','15 per class; 35 English, 35 Chinese and 35 mixed-language requests.'],
    ],[1.30,1.18,4.19],numeric=(1,))
    p(d,'The final synthetic test was authored separately without inspecting the training examples and frozen before evaluation. Feature extraction was fitted only on the training set. Six fixed candidate settings compared character-only features with a character-and-word combination at C values of 1, 4 and 12. Selection used validation macro F1, with accuracy as a tie-breaker. The selected model achieved 74/98 validation accuracy, or 75.5%.')
    p(d,'The final test was evaluated after the model and decision policy were frozen. Its errors were retained in the report rather than used to retune this version. Normalised full-text duplicates and shared family identifiers were checked across splits. This reduces exact leakage but cannot remove shared vocabulary, authoring style or semantic-template bias. A new untouched test set is needed if these final-test errors inform a later model revision [3].')
    h(d,'Handling uncertainty')
    p(d,'A direct supported-focus decision requires a raw top-class score of at least 0.50, a top-two score margin of at least 0.05, and a known-feature fraction of at least 0.15. The last measure indicates how much of the request matches the learned character-feature vocabulary. The score and margin settings were selected on validation data; the known-feature floor was fixed at 0.15. All three remained unchanged during integration.')
    p(d,'The class scores are uncalibrated softmax values. They rank the model outputs but are not validated probabilities that a classification is correct. A low score, a narrow margin or an unfamiliar input prompts clarification. Unsupported outputs are blocked. This is a limited uncertainty mechanism, not comprehensive detection of unsafe or out-of-domain requests.')

    page(d,'Predictive results and interpretation')
    p(d,'The following results come from the frozen set of 105 synthetic requests. Raw accuracy measures the highest-scoring class before the uncertainty policy. Macro F1 gives each of the seven classes equal weight. Neither metric measures real-member outcomes or the quality of an entire workout plan.')
    table(d,['Measure','Observed result','Interpretation'],[
        ['Raw seven-class accuracy','76/105  |  72.4%','Correct highest-scoring label before clarification.'],
        ['Raw macro F1','0.722','Average F1 across the seven classes.'],
        ['Direct focus coverage','42/105  |  40.0%','Requests accepted as one of the five training focuses.'],
        ['Precision of direct focuses','36/42  |  85.7%','Correct among accepted focuses; not overall accuracy.'],
        ['Final-label accuracy','61/105  |  58.1%','Counts clarification as the policy output.'],
        ['Unsupported requests accepted','1 request','The rejection mechanism is not a reliable safety filter.'],
    ],[1.65,1.48,3.54],numeric=(1,))
    p(d,'Coverage and precision must be reported together. The uncertainty policy reduces the number of direct decisions, but it does not establish an improvement in overall classification accuracy. Its 85.7% accepted-focus precision on the final test was below the 90% target used during validation selection. The integrated user-confirmation step has not been evaluated in a human study, so no accuracy after user confirmation is claimed.')
    table(d,['Language','Correct requests','Raw accuracy'],[
        ['English','27/35','77.1%'],['Chinese','25/35','71.4%'],['Mixed language','24/35','68.6%'],
    ],[2.7,1.8,2.17],numeric=(1,2))
    p(d,'The lower mixed-language result indicates a useful priority for further data collection, although each subgroup contains only 35 constructed requests. Upper-body recall was 8/15, or 53.3%, and was the weakest class recall in this test. The data are too small and artificial to establish population-level differences between languages or member groups.')
    h(d,'Size and runtime evidence')
    p(d,'The JSON model is 1,424,885 bytes. On the development machine, Python inference had a median of 0.202 ms and a 95th-percentile time of 0.620 ms across the test requests. These measurements exclude model loading and downloading. They are not browser or mobile-device benchmarks. Reproduction produced identical feature ordering, TF-IDF values, weights, biases and uncertainty policy; timestamps and elapsed-time metadata caused the complete file hashes to differ.')

    page(d,'Demo integration and software checks')
    p(d,'The released demo loads the frozen model and runs classification on the visitor\'s device. A high-confidence supported focus creates a validated intent. An uncertain request offers focus choices while retaining rule-extracted time, preferences and exclusions. The selected focus is attributed to the user, rather than reported as successful model recognition. Editing the request invalidates a pending confirmation. Medical, invalid-time and unsupported-experience checks run before confirmation and cannot be bypassed by choosing a focus.')
    p(d,'The interface then passes the intent to the existing scheduler. The entrance-based floor plan, equipment states, aisle routes, session timeline and full simulation remain functional. Synthetic events include another member extending use, the user taking extra rest, and a later station becoming occupied. They change actual scheduler state and the remaining forecast while preserving an exercise already in progress.')
    table(d,['Integration test area','Passed','Evidence covered'],[
        ['UI and offline artifact','17','Local classification, confirmation, stale text, controls and bundled operation.'],
        ['Server boundary','7','Public assets, retired endpoint, hidden paths and no provider calls.'],
        ['Scheduling engine','16','Active users, conflicts, availability provenance and feasible suggestions.'],
        ['Field extraction','13','Durations, negation, exclusions, preferences and validation.'],
        ['Local model adapter','11','Frozen provenance, learned focus, guarded confirmation and no network dependency.'],
        ['Forecast and simulation','14','Timed routes, disruptions, changed order and deterministic playback.'],
        ['Total','78','Functional checks; separate from prediction accuracy.'],
    ],[1.65,.75,4.27],numeric=(1,))
    p(d,'The standalone model also passed eight software checks: six runtime checks, including Python and JavaScript parity across 20 separately written Unicode and boundary probes, plus two offline-page checks. Per-class scores agreed within 1e-7. These tests establish consistency of implementation under the tested inputs; they do not increase the measured 72.4% classification accuracy.')
    p(d,'The public release was verified to return the page and matching model asset without a login redirect. Its health response identifies the local classifier and reports no external API calls. The previous planning endpoint now returns HTTP 410, and the active server does not read model credentials. A first website visit downloads assets; the self-contained HTML embeds them for server-independent use.')
    p(d,'Browser visual and responsive inspection remains incomplete because browser access was blocked. The DOM harness does not establish animation smoothness, map legibility, browser-specific accessibility or phone performance.')

    page(d,'Limitations and further validation')
    p(d,'The present evidence supports an experimental classroom prototype. It does not establish reliable understanding of arbitrary member requests, safe exercise selection, reduced waiting, commercial value or real-time gym operation. The limited synthetic corpus and heterogeneous unsupported class are the main model weaknesses. Existing rule checks offer additional boundaries, but a false accepted unsupported request remains evidence that this system cannot provide medical screening.')
    p(d,'The gym state and usage durations are simulated independently in each browser. There is no live sensor feed, shared occupancy database or trained duration predictor. Forecast release times are conditional estimates; confirming a future step does not reserve equipment. Exercise-substitution relationships still require qualified trainer review before a real-gym pilot.')
    h(d,'Proposed next validation')
    p(d,'First, collect consented member requests and obtain independent labels. Reserve a new test set before revising the model, covering negation, corrections, mixed language, missing goals and unsupported restrictions. Compare the trained classifier with the existing rule-based parser on the same requests. Report raw accuracy, macro F1, direct-decision coverage, accepted-focus precision and unsupported false acceptance together.')
    p(d,'Second, test the complete interface on agreed desktop and mobile devices. Measure loading and response time separately, and observe whether users identify the next station, notice a changed route and understand confirmation versus reservation. Measure confirmation completion and correction errors directly rather than inferring them from software tests.')
    p(d,'Third, validate equipment-state accuracy and usage-duration estimates against observed finishes, then compare fixed-order and adaptive scheduling under matched demand. Report waiting, completion, route changes and fairness separately. Real-gym benefit should be claimed only after those measurements, with reviewed exercise rules and suitable operating controls.')
    h(d,'Evidence records')
    p(d,'Model evidence: NextSet Flow Intent Model/reports/metrics.json; test_predictions.jsonl; reproducibility.json; and model/selection.json. Integration evidence: NextSet Flow Demo/tests/*.test.mjs; dist/local-model-info.json; and Technical Validation v4.md. The original seed, validation and frozen final-test files remain available for audit.',size=9)
    p(d,'Frozen model SHA-256',size=9)
    hash_p=p(d,METRICS['modelSHA256'],size=8.8)
    for run in hash_p.runs:run.font.name='Consolas'
    p(d,'Demo release source commit 385ad22ab378830a4a3f6c6e452707760f789cf9',size=9)
    link(d,'Public demo  nextset-flow-demo.lilacghost32768.chatgpt.site','https://nextset-flow-demo.lilacghost32768.chatgpt.site')
    h(d,'Method references')
    link(d,'[1] scikit-learn 1.9.1  TfidfVectorizer  accessed 27 October 2026','https://scikit-learn.org/stable/modules/generated/sklearn.feature_extraction.text.TfidfVectorizer.html')
    link(d,'[2] scikit-learn 1.9.1  LogisticRegression  accessed 27 October 2026','https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LogisticRegression.html')
    link(d,'[3] scikit-learn 1.9.1  Common pitfalls and data leakage  accessed 27 October 2026','https://scikit-learn.org/stable/common_pitfalls.html#data-leakage')
    d.save(OUTPUT)
    print(OUTPUT)
    print('Paragraphs',len(d.paragraphs),'Tables',len(d.tables),'Words',sum(len(x.text.split()) for x in d.paragraphs)+sum(len(c.text.split()) for t in d.tables for row in t.rows for c in row.cells))

if __name__=='__main__':make()
